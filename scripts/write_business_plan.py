#!/usr/bin/env python3
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

SOURCE = Path("/Users/safabadamchi/Desktop/Experience/Work/BBLS/BBLS_Business_Plan_Brand_Launch_Final.docx")
OUTPUT = Path("/Users/safabadamchi/Desktop/Experience/Work/BBLS/BBLS_Business_Plan_Website_Design_Studio.docx")

NAVY = RGBColor(0x10, 0x18, 0x20)
BLUE = RGBColor(0x2E, 0x5E, 0x8C)
GOLD = RGBColor(0xB2, 0x8A, 0x55)
BODY = RGBColor(0x20, 0x25, 0x2A)
MUTED = RGBColor(0x68, 0x70, 0x78)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_run(run, text, *, size=11, bold=False, italic=False, color=BODY, name="Aptos"):
    run.text = text
    run.bold = bold
    run.italic = italic
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)


def add_para(doc, text, *, size=11, bold=False, italic=False, color=BODY, space_after=7, space_before=0, align=None, name="Aptos"):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(space_after)
    para.paragraph_format.space_before = Pt(space_before)
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if align:
        para.alignment = align
    set_run(para.add_run(text), text, size=size, bold=bold, italic=italic, color=color, name=name)
    return para


def add_h1(doc, text):
    para = doc.add_paragraph(text, style="Heading 1")
    para.paragraph_format.space_before = Pt(12)
    para.paragraph_format.space_after = Pt(6)
    para.paragraph_format.keep_with_next = True
    if para.runs:
        set_run(para.runs[0], text, size=17, bold=True, color=NAVY, name="Aptos Display")
    return para


def add_h2(doc, text):
    para = doc.add_paragraph(text, style="Heading 2")
    if para.runs:
        set_run(para.runs[0], text, size=13, bold=True, color=BLUE, name="Aptos Display")
    return para


def add_h3(doc, text):
    para = doc.add_paragraph(text, style="Heading 3")
    if para.runs:
        set_run(para.runs[0], text, size=12, bold=True, color=NAVY, name="Aptos Display")
    return para


def add_label(doc, text):
    return add_para(doc, text, size=10, bold=True, color=GOLD, space_after=4, space_before=6)


def add_body(doc, text):
    return add_para(doc, text, size=11, color=BODY, space_after=7)


def add_bullet(doc, text):
    para = doc.add_paragraph(text, style="List Bullet")
    para.paragraph_format.space_after = Pt(3)
    if para.runs:
        set_run(para.runs[0], text, size=10.5, color=BODY)
    else:
        set_run(para.add_run(text), text, size=10.5, color=BODY)
    return para


def shade(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "D8DEE4")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def fill_cell(cell, text, *, header=False):
    cell.text = ""
    para = cell.paragraphs[0]
    para.paragraph_format.space_before = Pt(3)
    para.paragraph_format.space_after = Pt(3)
    run = para.add_run(text)
    set_run(run, text, size=9, bold=header, color=WHITE if header else BODY)
    shade(cell, "101820" if header else "FFFFFF")
    set_cell_border(cell)
    cell.width = Inches(1.6)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.allow_autofit = False
    for i, header in enumerate(headers):
        fill_cell(table.rows[0].cells[i], header, header=True)
        table.columns[i].width = Inches(widths[i])
        table.rows[0].cells[i].width = Inches(widths[i])
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            fill_cell(table.rows[r].cells[c], value)
            table.rows[r].cells[c].width = Inches(widths[c])
    return table


def clear_body(doc):
    body = doc.element.body
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        body.remove(child)


def copy_footer(source, dest):
    dest.sections[0].footer.is_linked_to_previous = False
    dest_para = dest.sections[0].footer.paragraphs[0]
    dest_para.clear()
    set_run(
        dest_para.add_run("BBLS — Boutique Brand & Launch Studio  |  Confidential planning document"),
        "BBLS — Boutique Brand & Launch Studio  |  Confidential planning document",
        size=9,
        color=MUTED,
    )


def build():
    source = Document(str(SOURCE))
    doc = Document(str(SOURCE))
    clear_body(doc)
    sec = doc.sections[0]
    src_sec = source.sections[0]
    sec.page_width = src_sec.page_width
    sec.page_height = src_sec.page_height
    sec.left_margin = src_sec.left_margin
    sec.right_margin = src_sec.right_margin
    sec.top_margin = src_sec.top_margin
    sec.bottom_margin = src_sec.bottom_margin
    copy_footer(source, doc)

    add_para(doc, "BUSINESS PLAN | 2026", size=9.5, bold=True, color=GOLD, space_after=10)
    add_para(doc, "BBLS", size=34, bold=True, color=NAVY, space_after=2)
    add_para(doc, "Boutique Brand & Launch Studio", size=18, color=BLUE, space_after=4)
    add_para(doc, "Premium Website Design Studio in Orange County", size=13, bold=True, color=NAVY, space_after=8)
    add_para(doc, "We design websites built to launch, convert, and grow.", size=13, italic=True, color=MUTED, space_after=10)
    add_para(doc, "WE DON’T SHAPE BODIES.\nWE BUILD BUSINESSES.", size=16, bold=True, color=NAVY, space_after=10)
    add_para(doc, "Prepared for internal planning, partner conversations, and studio execution", size=9.5, color=MUTED, space_after=2)
    add_para(doc, "Orange County, California | bbls.studio", size=9.5, bold=True, color=BLUE, space_after=16)

    add_h1(doc, "1. Executive Summary")
    add_body(
        doc,
        "BBLS — Boutique Brand & Launch Studio is a premium website design and launch studio in Orange County. BBLS helps businesses launch polished, high-performing websites designed to build credibility, generate leads, and support growth.",
    )
    add_label(doc, "Primary service message")
    add_para(doc, "We design websites built to launch, convert, and grow.", size=14, bold=True, color=NAVY, space_after=8)
    add_body(
        doc,
        "The studio offers three primary website packages: Simple Landing Page, starting at $3,000; Business Website, starting at $5,500; and Commercial Website, starting at $8,500. Brand Identity is available only as an optional add-on, starting at $3,500. Work beyond ten pages, or work that requires e-commerce, booking systems, memberships, multilingual sites, client portals, or advanced integrations, is scoped after consultation.",
    )
    add_h2(doc, "Vision and Mission")
    add_body(doc, "Vision: Become a recognized premium website design studio in Orange County for businesses that want a polished site built to launch, convert, and grow.")
    add_body(doc, "Mission: Design and launch websites that present the business clearly, support inquiries, and give the company a credible commercial presence.")
    add_h2(doc, "Strategic Objectives")
    add_bullet(doc, "Establish BBLS as a premium website design and launch studio in Orange County.")
    add_bullet(doc, "Sell from a clear three-package website offer, with Brand Identity available only as an add-on.")
    add_bullet(doc, "Keep project scopes written, priced as starting points, and confirmed before production begins.")
    add_bullet(doc, "Use selected live work as proof, without invented results or unverified claims.")
    add_h2(doc, "Value Proposition")
    add_body(doc, "BBLS gives a business one studio for website strategy, design, build, and launch. Clients do not need to coordinate a separate designer, developer, and launch resource for a standard commercial site.")

    add_h1(doc, "2. Company Description")
    add_body(doc, "BBLS is an Orange County studio focused on premium website design and launch. The work is design-led, commercial, and direct. The studio plans the site around the offer, the audience, and the action the business needs the visitor to take.")
    add_body(doc, "BBLS is not a general marketing agency or a software product company. This plan covers project-based website design and optional brand work only.")
    add_h2(doc, "Studio Identity")
    add_bullet(doc, "Name: BBLS — Boutique Brand & Launch Studio")
    add_bullet(doc, "Positioning line: Premium Website Design Studio in Orange County")
    add_bullet(doc, "Website: bbls.studio")
    add_bullet(doc, "Email: hello@bbls.studio")
    add_bullet(doc, "Phone: (949) 524-2324")
    add_h2(doc, "Brand Statement")
    add_para(doc, "We don’t shape bodies. We build businesses.", size=13, bold=True, color=NAVY, space_after=8)
    add_body(doc, "The line marks BBLS as a commercial studio. The work is about identity, websites, and launch — not fitness, wellness, or body-focused brands as a category.")

    add_h1(doc, "3. Positioning")
    add_body(doc, "BBLS is positioned first as a premium website design and launch studio. Brand work supports the website when a business needs a clearer identity before or during the build. Brand Identity is not sold as a stand-alone primary package in this plan.")
    add_h2(doc, "Market Position")
    add_body(doc, "BBLS sits between low-cost template vendors and large agencies. The studio competes on focused scope, visual quality, conversion-minded structure, and a defined starting price. It does not compete on the lowest bid.")
    add_h2(doc, "What BBLS Sells")
    add_bullet(doc, "Simple Landing Page — one focused commercial page.")
    add_bullet(doc, "Business Website — a complete site of up to five pages.")
    add_bullet(doc, "Commercial Website — a larger conversion-focused site of six to ten pages.")
    add_bullet(doc, "Optional Brand Identity add-on when the website needs a usable visual system.")
    add_h2(doc, "What This Plan Does Not Include")
    add_bullet(doc, "Recurring monthly maintenance as a published offer.")
    add_body(doc, "This plan covers project website design and optional brand add-ons only. Custom requirements are handled through consultation, not through a fourth primary package.")

    add_h1(doc, "4. Target Market")
    add_body(doc, "BBLS works with founders and established businesses that need a polished website to look credible, collect inquiries, and support growth.")
    add_h2(doc, "Primary Clients")
    add_bullet(doc, "Founders and small businesses that need one focused, professional page.")
    add_bullet(doc, "Established businesses that need a complete and credible online presence.")
    add_bullet(doc, "Growth-focused companies that need a larger, conversion-focused website.")
    add_h2(doc, "Fit Criteria")
    add_bullet(doc, "A defined offer, service, or product the website can present.")
    add_bullet(doc, "Willingness to provide content, access, and timely approvals.")
    add_bullet(doc, "A need for a custom-designed site rather than an unedited template.")
    add_h2(doc, "Geographic Focus")
    add_body(doc, "The studio is based in Orange County and can work with clients beyond the region. Local presence supports in-person conversations when useful. Delivery does not depend on the client being local.")

    add_h1(doc, "5. Services")
    add_body(doc, "BBLS has exactly three primary website packages. Each package is a starting point. Final pricing depends on page count, content requirements, integrations, functionality, timeline, and project complexity.")
    add_h2(doc, "Simple Landing Page — Starting at $3,000")
    add_body(doc, "For founders and small businesses that need one focused, professional page.")
    add_bullet(doc, "One custom-designed landing page")
    add_bullet(doc, "Mobile and desktop responsive design")
    add_bullet(doc, "Clear content hierarchy and call-to-action strategy")
    add_bullet(doc, "Contact or lead-generation form")
    add_bullet(doc, "Basic on-page SEO setup")
    add_bullet(doc, "Analytics integration")
    add_bullet(doc, "Launch support")
    add_bullet(doc, "Up to two revision rounds")
    add_h2(doc, "Business Website — Starting at $5,500")
    add_body(doc, "For established businesses that need a complete and credible online presence.")
    add_bullet(doc, "Up to five custom-designed pages")
    add_bullet(doc, "Responsive website design")
    add_bullet(doc, "Strategic user experience and navigation")
    add_bullet(doc, "Contact, inquiry, or lead-generation forms")
    add_bullet(doc, "Basic CMS setup where appropriate")
    add_bullet(doc, "Basic on-page SEO")
    add_bullet(doc, "Analytics integration")
    add_bullet(doc, "Launch support")
    add_bullet(doc, "Up to two revision rounds")
    add_h2(doc, "Commercial Website — Starting at $8,500")
    add_body(doc, "For growth-focused companies that need a larger, conversion-focused website.")
    add_bullet(doc, "Six to ten custom-designed pages")
    add_bullet(doc, "Custom UX and website architecture")
    add_bullet(doc, "Conversion-focused page layouts")
    add_bullet(doc, "Advanced forms or third-party integrations")
    add_bullet(doc, "CMS setup where appropriate")
    add_bullet(doc, "On-page SEO foundation")
    add_bullet(doc, "Analytics and conversion tracking")
    add_bullet(doc, "Launch support")
    add_bullet(doc, "Up to three revision rounds")
    add_h2(doc, "Custom Requirements")
    add_body(doc, "For e-commerce, booking systems, memberships, multilingual websites, client portals, advanced integrations, or projects larger than ten pages:")
    add_para(doc, "Custom scope and pricing available after consultation.", size=12, bold=True, color=NAVY, space_after=8)
    add_h2(doc, "Brand Identity Add-on — Starting at $3,500")
    add_body(doc, "Brand Identity is not a primary package. It is an optional website add-on when the project needs a usable visual system.")
    add_bullet(doc, "Logo system")
    add_bullet(doc, "Color palette")
    add_bullet(doc, "Typography system")
    add_bullet(doc, "Core visual direction")
    add_bullet(doc, "Basic brand guidelines")

    add_h1(doc, "6. Pricing")
    add_body(doc, "All published prices use “Starting at” language. They are planning figures, not fixed quotes. Final pricing depends on page count, content requirements, integrations, functionality, timeline, and project complexity. Third-party services and platform costs are quoted separately when required. Final scope and price are confirmed in writing before work begins.")
    add_table(
        doc,
        ["Offer", "Starting price", "Scope", "Best for"],
        [
            ["Simple Landing Page", "Starting at $3,000", "One custom page; responsive design; CTA strategy; form; basic SEO; analytics; launch; 2 revision rounds", "Founders and small businesses that need one focused page"],
            ["Business Website", "Starting at $5,500", "Up to 5 pages; UX and navigation; forms; basic CMS where appropriate; SEO; analytics; launch; 2 revision rounds", "Established businesses that need a complete presence"],
            ["Commercial Website", "Starting at $8,500", "6–10 pages; custom UX; conversion layouts; advanced forms or integrations; CMS; SEO; tracking; launch; 3 revision rounds", "Growth-focused companies that need a larger site"],
            ["Brand Identity Add-on", "Starting at $3,500", "Logo, color, type, visual direction, and basic guidelines", "Website projects that also need a usable identity system"],
            ["Custom work", "Quoted after consultation", "E-commerce, booking, memberships, multilingual sites, portals, advanced integrations, or more than 10 pages", "Projects outside the three primary packages"],
        ],
        [1.55, 1.25, 2.35, 1.65],
    )

    add_h1(doc, "7. Delivery Model")
    add_body(doc, "BBLS uses a four-stage process. Timing depends on scope, content readiness, access, and approvals.")
    add_h3(doc, "01 Discover")
    add_body(doc, "Clarify the offer, audience, goals, page count, required materials, and launch target.")
    add_h3(doc, "02 Define")
    add_body(doc, "Set the scope, site structure, content needs, and creative direction before production starts.")
    add_h3(doc, "03 Design")
    add_body(doc, "Design the pages, hierarchy, and key flows. Add brand work only when the Brand Identity add-on is included.")
    add_h3(doc, "04 Deliver")
    add_body(doc, "Build, review, connect forms and analytics, launch the site, and hand off the work.")
    add_label(doc, "Timing note")
    add_body(doc, "Focused landing-page launches may finish in 14 days when scope, content, and approvals are ready. Broader website work typically takes two to four weeks. Larger or custom work is scheduled after the scope is confirmed.")

    add_h1(doc, "8. Competitive Advantage")
    add_bullet(doc, "A clear three-package website offer with starting prices a client can understand.")
    add_bullet(doc, "Design and launch handled in one studio, so the first impression and the live site stay aligned.")
    add_bullet(doc, "Conversion-minded structure: hierarchy, calls to action, forms, and analytics are part of the package, not afterthoughts.")
    add_bullet(doc, "Orange County presence with a focused commercial offer, not a generic service list.")
    add_bullet(doc, "Optional brand add-on when the website needs identity work, without turning brand into a separate primary product.")
    add_body(doc, "BBLS does not claim awards, exclusive partnerships, guaranteed rankings, guaranteed traffic, or guaranteed conversions. Advantage is the work itself: a polished site, a defined scope, and a direct process.")

    add_h1(doc, "9. Sales Strategy")
    add_body(doc, "Sales begin with a focused conversation: what the business is building, what the website needs to do, and which package is the closest fit.")
    add_h2(doc, "Lead Sources")
    add_bullet(doc, "bbls.studio, with website design in Orange County as the primary public message.")
    add_bullet(doc, "Direct outreach and referrals in Orange County and the greater Southern California market.")
    add_bullet(doc, "Inbound email to hello@bbls.studio and calls to (949) 524-2324.")
    add_h2(doc, "Sales Motion")
    add_bullet(doc, "Qualify page count, content readiness, integrations, and timeline.")
    add_bullet(doc, "Map the work to Simple Landing Page, Business Website, or Commercial Website.")
    add_bullet(doc, "Offer Brand Identity only when the website needs a visual system.")
    add_bullet(doc, "Move custom or larger work to consultation instead of stretching a package.")
    add_bullet(doc, "Confirm scope, starting price, and payment terms in writing before kickoff.")
    add_h2(doc, "Proof")
    add_body(doc, "Public proof is limited to live work the studio can show: Raysan IP, 7 Stud Farm, and IP Law Nerds. This plan does not assign results, conversion rates, or revenue to those projects.")

    add_h1(doc, "10. Revenue Model")
    add_body(doc, "Revenue in this plan comes from project work only.")
    add_bullet(doc, "Simple Landing Page — starting at $3,000")
    add_bullet(doc, "Business Website — starting at $5,500")
    add_bullet(doc, "Commercial Website — starting at $8,500")
    add_bullet(doc, "Optional Brand Identity add-on — starting at $3,500")
    add_bullet(doc, "Custom website work — priced after consultation")
    add_body(doc, "This version of the plan does not include recurring monthly-maintenance revenue.")
    add_h2(doc, "Pricing Controls")
    add_bullet(doc, "Use deposits and milestones tied to discovery, design approval, and launch.")
    add_bullet(doc, "Use written change orders when page count, integrations, or functionality expand.")
    add_bullet(doc, "Quote third-party tools and platform costs separately.")

    add_h1(doc, "11. Financial Projections")
    add_body(doc, "The figures below are illustrative planning examples. They use only the starting prices in this plan. They are not forecasts, targets promised to clients, or statements of past revenue.")
    add_h2(doc, "Planning Assumptions")
    add_bullet(doc, "Simple Landing Page counted at $3,000.")
    add_bullet(doc, "Business Website counted at $5,500.")
    add_bullet(doc, "Commercial Website counted at $8,500.")
    add_bullet(doc, "Brand Identity add-on counted at $3,500 when included.")
    add_bullet(doc, "No recurring-support revenue is included.")
    add_bullet(doc, "Actual invoices may be higher after scope is confirmed.")
    add_h2(doc, "Illustrative Monthly Mixes")
    add_table(
        doc,
        ["Scenario", "Illustrative mix", "Websites", "Brand add-ons", "Total"],
        [
            ["Foundation", "1 Simple Landing Page; 1 Business Website", "$8,500", "$0", "$8,500"],
            ["Target", "1 Simple Landing Page; 1 Business Website; 1 Commercial Website", "$17,000", "$0", "$17,000"],
            ["Growth", "1 Business Website; 2 Commercial Websites; 2 Brand Identity add-ons", "$22,500", "$7,000", "$29,500"],
        ],
        [1.15, 2.55, 1.05, 1.15, 0.9],
    )
    add_body(doc, "Check: $3,000 + $5,500 = $8,500. $3,000 + $5,500 + $8,500 = $17,000. $5,500 + $8,500 + $8,500 + $3,500 + $3,500 = $29,500.")
    add_h2(doc, "Illustrative Annual View")
    add_body(doc, "If a planning year contained twelve Target-mix months, project revenue at starting prices would be $204,000. If a planning year contained six Foundation months and six Target months, project revenue at starting prices would be $153,000. These are capacity illustrations only. They do not represent booked work or historical results.")
    add_h2(doc, "Operating Priorities")
    add_bullet(doc, "Protect margin by qualifying content readiness and limiting concurrent builds.")
    add_bullet(doc, "Track hours, package mix, close rate, and add-on attachment.")
    add_bullet(doc, "Keep Brand Identity attached to website work rather than sold as a stand-alone primary offer.")

    add_h1(doc, "12. Website and Marketing Copy")
    add_body(doc, "Public and proposal language should stay consistent with this plan.")
    add_h2(doc, "Identity Lines")
    add_bullet(doc, "BBLS — Boutique Brand & Launch Studio")
    add_bullet(doc, "Premium Website Design Studio in Orange County")
    add_h2(doc, "Primary Service Message")
    add_para(doc, "We design websites built to launch, convert, and grow.", size=13, bold=True, color=NAVY, space_after=8)
    add_h2(doc, "Brand Statement")
    add_para(doc, "We don’t shape bodies. We build businesses.", size=13, bold=True, color=NAVY, space_after=8)
    add_h2(doc, "Supporting Copy")
    add_body(doc, "BBLS helps businesses launch polished, high-performing websites designed to build credibility, generate leads, and support growth.")
    add_h2(doc, "Package Lines")
    add_bullet(doc, "Simple Landing Page, starting at $3,000. One focused, professional page.")
    add_bullet(doc, "Business Website, starting at $5,500. A complete and credible online presence.")
    add_bullet(doc, "Commercial Website, starting at $8,500. A larger, conversion-focused website.")
    add_bullet(doc, "Brand Identity add-on, starting at $3,500. Optional, and used when the website needs a visual system.")
    add_bullet(doc, "Custom scope and pricing available after consultation.")
    add_h2(doc, "Contact Lines")
    add_bullet(doc, "Ready to build a more distinctive business?")
    add_bullet(doc, "Tell BBLS what you are building, what you need, and where you want the business to go.")
    add_bullet(doc, "hello@bbls.studio")
    add_bullet(doc, "(949) 524-2324")

    add_h1(doc, "13. 90-Day Action Plan")
    add_bullet(doc, "Publish the three website packages, the Brand Identity add-on, and the custom-scope line in proposals and on bbls.studio.")
    add_bullet(doc, "Remove any remaining outdated package language from public and internal materials.")
    add_bullet(doc, "Standardize discovery notes, written scopes, deposits, and revision limits.")
    add_bullet(doc, "Keep live work current: Raysan IP, 7 Stud Farm, and IP Law Nerds.")
    add_bullet(doc, "Run a focused Orange County outreach plan around website design and launch.")
    add_bullet(doc, "Review package mix, utilization, and margin after the first five paid website engagements.")
    add_label(doc, "Positioning line")
    last = add_para(doc, "We don’t shape bodies. We build businesses.", size=14, bold=True, color=NAVY, space_after=0, space_before=0)
    last.paragraph_format.keep_together = True

    doc.save(str(OUTPUT))
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    build()
